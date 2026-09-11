"""Собрать Excel с исходными таблицами всех графиков и источниками."""

import argparse
import csv
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent


def convert(value):
    if value in ("True", "False"):
        return value == "True"
    try:
        return int(value)
    except ValueError:
        try:
            return float(value)
        except ValueError:
            return value


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, default=ROOT / "documentation")
    args = parser.parse_args()
    args.out.mkdir(exist_ok=True, parents=True)
    workbook = Workbook()
    cover = workbook.active
    cover.title = "Как читать"
    for row in [
        ["Эксперименты по лабораторным 1–8", "Вариант 15, seed 45"],
        [
            "seconds",
            "Медиана пяти повторов, секунды; в пакетных опытах — на одну операцию",
        ],
        ["trial_1 … trial_5", "Исходные времена пяти повторов в тех же единицах"],
        ["comparisons", "Сравнения ключей или символов; в ЛР 6 — среднее на запрос"],
        ["Подготовка", "Один прогрев, ввод и подготовка свежих состояний вне таймера"],
        ["Ограничения", "Фоновые процессы и питание не контролировались"],
        [
            "ЛР 8 alpha",
            "Для dict — внешняя нагрузка n/1024, не его внутренняя загрузка",
        ],
        [
            "Графики",
            "PNG и объяснения находятся в отчётах; каждый лист соответствует исходному CSV",
        ],
    ]:
        cover.append(row)
    sources = workbook.create_sheet("Источники")
    sources.append(
        ["ЛР", "CSV в репозитории", "CPU", "Python", "Дата UTC", "Вариант", "Seed"]
    )
    expected = {}
    for number in range(1, 9):
        directory = ROOT / "labs" / f"lab{number:02}" / "results"
        environment = json.loads(
            (directory / "environment.json").read_text(encoding="utf-8")
        )
        for path in sorted(directory.glob("*.csv")):
            title = f"{number:02}_{path.stem}"[:31]
            sheet = workbook.create_sheet(title)
            with path.open(encoding="utf-8", newline="") as stream:
                rows = list(csv.reader(stream))
            for index, row in enumerate(rows):
                sheet.append(row if index == 0 else [convert(value) for value in row])
            expected[title] = len(rows)
            url = (
                "https://github.com/potemkinmihail777-dot/data-structures-and-algorithms-variant-15/blob/main/"
                + path.relative_to(ROOT).as_posix()
            )
            sources.append(
                [
                    number,
                    url,
                    environment["cpu"],
                    environment["python"],
                    environment["timestamp_utc"],
                    15,
                    45,
                ]
            )
            sources.cell(sources.max_row, 2).hyperlink = url
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="33465B")
        for column in sheet.columns:
            index = column[0].column
            width = min(
                60, max(13, max(len(str(cell.value or "")) for cell in column) + 2)
            )
            sheet.column_dimensions[get_column_letter(index)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, float):
                    cell.number_format = (
                        "0.000000000"
                        if "trial" in str(sheet.cell(1, cell.column).value)
                        or sheet.cell(1, cell.column).value == "seconds"
                        else "0.0000"
                    )
    target = args.out / "experiments_variant15.xlsx"
    workbook.save(target)
    checked = load_workbook(target, read_only=True, data_only=True)
    assert all(checked[name].max_row == count for name, count in expected.items())
    print(
        f"Workbook: {len(checked.sheetnames)} sheets, {sum(expected.values())} source rows including headers"
    )


if __name__ == "__main__":
    main()
