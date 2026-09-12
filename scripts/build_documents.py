"""Создать восемь отдельных учебных документов в PDF и редактируемом DOCX."""

import argparse
from html import escape
from pathlib import Path
import re
from urllib.parse import quote
import zipfile
import markdown
from lxml import html
import matplotlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Image,
    Table,
    TableStyle,
    PageBreak,
)

ROOT = Path(__file__).resolve().parent.parent
GITHUB = "https://github.com/potemkinmihail777-dot/data-structures-and-algorithms-variant-15/blob/main/"


def target_url(href, folder):
    if re.match(r"https?://", href):
        return href
    resolved = (folder / href).resolve()
    return GITHUB + quote(resolved.relative_to(ROOT).as_posix())


def markup(element, folder):
    result = escape(element.text or "")
    for child in element:
        inside = markup(child, folder)
        if child.tag in ("strong", "b"):
            result += "<b>" + inside + "</b>"
        elif child.tag in ("em", "i"):
            result += "<i>" + inside + "</i>"
        elif child.tag == "code":
            result += '<font name="DejaVuMono">' + inside + "</font>"
        elif child.tag == "a":
            result += (
                '<a href="'
                + escape(target_url(child.get("href"), folder), quote=True)
                + '">'
                + inside
                + "</a>"
            )
        elif child.tag == "br":
            result += "<br/>"
        elif child.tag != "img":
            result += inside
        result += escape(child.tail or "")
    return result


def word_inline(paragraph, element, folder):
    if element.text:
        paragraph.add_run(element.text.replace("\n", " "))
    for child in element:
        content = "".join(child.itertext())
        if child.tag == "a":
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(
                qn("r:id"),
                paragraph.part.relate_to(
                    target_url(child.get("href"), folder),
                    RT.HYPERLINK,
                    is_external=True,
                ),
            )
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "245B84")
            properties.append(color)
            run.append(properties)
            text = OxmlElement("w:t")
            text.text = content
            run.append(text)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
        elif child.tag in ("p", "span"):
            word_inline(paragraph, child, folder)
        elif child.tag != "img":
            run = paragraph.add_run(content)
            run.bold = child.tag in ("strong", "b")
            run.italic = child.tag in ("em", "i")
            if child.tag == "code":
                run.font.name = "Consolas"
                run.font.size = Pt(9)
        if child.tail:
            paragraph.add_run(child.tail.replace("\n", " "))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, default=ROOT / "documentation")
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    fonts = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
    for name, filename in [
        ("DejaVu", "DejaVuSans.ttf"),
        ("DejaVuBold", "DejaVuSans-Bold.ttf"),
        ("DejaVuItalic", "DejaVuSans-Oblique.ttf"),
        ("DejaVuMono", "DejaVuSansMono.ttf"),
    ]:
        pdfmetrics.registerFont(TTFont(name, str(fonts / filename)))
    pdfmetrics.registerFontFamily(
        "DejaVu",
        normal="DejaVu",
        bold="DejaVuBold",
        italic="DejaVuItalic",
        boldItalic="DejaVuBold",
    )
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="BodyRU",
            fontName="DejaVu",
            fontSize=10.2,
            leading=14.8,
            spaceAfter=7,
            alignment=TA_LEFT,
        )
    )
    styles.add(
        ParagraphStyle(
            name="TitleRU",
            parent=styles["BodyRU"],
            fontName="DejaVuBold",
            fontSize=20,
            leading=25,
            spaceAfter=17,
            keepWithNext=True,
        )
    )
    styles.add(
        ParagraphStyle(
            name="HeadingRU",
            parent=styles["BodyRU"],
            fontName="DejaVuBold",
            fontSize=13,
            leading=17,
            spaceBefore=12,
            spaceAfter=8,
            keepWithNext=True,
        )
    )
    styles.add(
        ParagraphStyle(
            name="CellRU",
            parent=styles["BodyRU"],
            fontSize=7.8,
            leading=10.6,
            spaceAfter=0,
        )
    )
    styles.add(
        ParagraphStyle(
            name="CodeRU",
            parent=styles["BodyRU"],
            fontName="DejaVuMono",
            fontSize=8.3,
            leading=12,
            leftIndent=8,
            spaceAfter=10,
        )
    )
    produced = []
    for number in range(1, 9):
        folder = ROOT / "labs" / f"lab{number:02}"
        doc = Document()
        section = doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(1.8)
        section.left_margin = section.right_margin = Cm(1.7)
        normal = doc.styles["Normal"]
        normal.font.name, normal.font.size = "Calibri", Pt(11)
        normal.paragraph_format.space_after = Pt(7)
        normal.paragraph_format.line_spacing = 1.12
        for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
            doc.styles[name].font.name = "Calibri"
            doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        for border in doc.styles.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
        footer = section.footer.paragraphs[0]
        footer.alignment = 1
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
        story = []
        for part_number, filename in enumerate(("notes.md", "report.md")):
            text = (folder / filename).read_text(encoding="utf-8")
            tree = html.fragment_fromstring(
                markdown.markdown(text, extensions=["tables", "fenced_code"]),
                create_parent="div",
            )
            if part_number:
                story.append(PageBreak())
            for element in tree:
                tag = element.tag
                if tag in ("h1", "h2", "h3"):
                    heading = doc.add_heading(
                        "".join(element.itertext()),
                        level=0 if tag == "h1" else int(tag[1]) - 1,
                    )
                    if part_number and tag == "h1":
                        heading.paragraph_format.page_break_before = True
                    story.append(
                        Paragraph(
                            markup(element, folder),
                            styles["TitleRU" if tag == "h1" else "HeadingRU"],
                        )
                    )
                elif tag == "p" and element.find("img") is not None:
                    for image in element.findall("img"):
                        path = folder / image.get("src")
                        doc.add_picture(str(path), width=Cm(17.1))
                        from PIL import Image as PILImage

                        with PILImage.open(path) as picture:
                            width, height = picture.size
                        displayed_width = A4[0] - 92
                        story += [
                            Image(
                                str(path),
                                width=displayed_width,
                                height=displayed_width * height / width,
                            ),
                            Spacer(1, 10),
                        ]
                elif tag == "p":
                    word_inline(doc.add_paragraph(), element, folder)
                    story.append(Paragraph(markup(element, folder), styles["BodyRU"]))
                elif tag in ("ul", "ol"):
                    for i, item in enumerate(element.findall("li"), 1):
                        word_inline(
                            doc.add_paragraph(
                                style="List Number" if tag == "ol" else "List Bullet"
                            ),
                            item,
                            folder,
                        )
                        prefix = f"{i}. " if tag == "ol" else "• "
                        story.append(
                            Paragraph(prefix + markup(item, folder), styles["BodyRU"])
                        )
                elif tag == "pre":
                    code = "".join(element.itertext())
                    p = doc.add_paragraph()
                    run = p.add_run(code)
                    run.font.name, run.font.size = "Consolas", Pt(9)
                    story.append(
                        Paragraph(
                            escape(code)
                            .replace("\n", "<br/>")
                            .replace("  ", "&#160;&#160;"),
                            styles["CodeRU"],
                        )
                    )
                elif tag == "table":
                    rows = element.findall(".//tr")
                    cells = [row.xpath("./th|./td") for row in rows]
                    columns = len(cells[0])
                    word_table = doc.add_table(rows=0, cols=columns)
                    word_table.style = "Table Grid"
                    pdf_rows = []
                    for row_index, row in enumerate(cells):
                        word_cells = word_table.add_row().cells
                        pdf_cells = []
                        for column, cell in enumerate(row):
                            word_inline(word_cells[column].paragraphs[0], cell, folder)
                            for run in word_cells[column].paragraphs[0].runs:
                                run.font.size = Pt(8)
                                if row_index == 0:
                                    run.bold = True
                            pdf_cells.append(
                                Paragraph(markup(cell, folder), styles["CellRU"])
                            )
                        pdf_rows.append(pdf_cells)
                        # Разрешаем перенос таблицы по строкам, но не разрыв строки.
                        tr_pr = word_table.rows[-1]._tr.get_or_add_trPr()
                        tr_pr.append(OxmlElement("w:cantSplit"))
                    header = OxmlElement("w:tblHeader")
                    word_table.rows[0]._tr.get_or_add_trPr().append(header)
                    available = A4[0] - 92
                    weights = [
                        max(
                            6,
                            min(
                                22,
                                max(len("".join(row[col].itertext())) for row in cells),
                            ),
                        )
                        for col in range(columns)
                    ]
                    widths = [available * w / sum(weights) for w in weights]
                    pdf_table = Table(
                        pdf_rows, colWidths=widths, repeatRows=1, hAlign="LEFT"
                    )
                    pdf_table.setStyle(
                        TableStyle(
                            [
                                (
                                    "BACKGROUND",
                                    (0, 0),
                                    (-1, 0),
                                    colors.HexColor("#E8ECF0"),
                                ),
                                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                (
                                    "GRID",
                                    (0, 0),
                                    (-1, -1),
                                    0.35,
                                    colors.HexColor("#BBC4CD"),
                                ),
                                ("TOPPADDING", (0, 0), (-1, -1), 5),
                                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                            ]
                        )
                    )
                    story += [pdf_table, Spacer(1, 10)]
                    doc.add_paragraph()
        stem = f"lab{number:02}_variant15"
        word_path, pdf_path = args.out / f"{stem}.docx", args.out / f"{stem}.pdf"
        doc.core_properties.title = f"Лабораторная {number}. Вариант 15"
        doc.core_properties.author = ""
        doc.save(word_path)

        def page_number(canvas, _document):
            canvas.saveState()
            canvas.setFont("DejaVu", 8)
            canvas.drawCentredString(A4[0] / 2, 24, str(canvas.getPageNumber()))
            canvas.restoreState()

        pdf = SimpleDocTemplate(
            str(pdf_path),
            pagesize=A4,
            leftMargin=46,
            rightMargin=46,
            topMargin=43,
            bottomMargin=43,
            title=f"Лабораторная {number}. Вариант 15",
            author="",
        )
        pdf.build(story, onFirstPage=page_number, onLaterPages=page_number)
        produced += [word_path, pdf_path]
        print("Documents", number, flush=True)
    with zipfile.ZipFile(
        args.out / "all_documentation_variant15.zip", "w", zipfile.ZIP_DEFLATED
    ) as archive:
        for file in produced:
            archive.write(file, file.name)


if __name__ == "__main__":
    main()
